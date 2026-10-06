"""日程任务业务逻辑：扁平表格的 CRUD、统计与 Excel 导出。

约定（本次改版）：
- 「一级计划」（level1）为分类列，「任务名」（title）是其下的具体任务；
- 不再有日 / 周 / 月度计划类型（月 / 周计划改由日历模块的独立文本承担）；
- 状态三态「未完成 / 进行中 / 已完成」完全由用户手动维护，不做基于日期的自动流转；
- 完成度为用户手填的 0-100 整数，与状态互不联动；
- 列表一律按结束日期升序排序（无结束日期的排最后）。
"""

import io
import json
import sqlite3
from typing import Optional

from .. import config, db
from ..models.planpool import (
    COMPLETION_MAX,
    COMPLETION_MIN,
    DEFAULT_COMPLETION,
    DEFAULT_PROGRESS,
    PROGRESS_DONE,
    PROGRESS_DOING,
    PROGRESS_OPTIONS,
    PROGRESS_TODO,
)

# 无结束日期的任务排在最后
ORDER_BY_SQL = "ORDER BY CASE WHEN end_date = '' THEN 1 ELSE 0 END, end_date, id"


def _get_conn() -> sqlite3.Connection:
    return db.get_conn(config.PLANPOOL_DB)


# ---------------------------------------------------------------------------
# 内置列种子 & 预设模板
# ---------------------------------------------------------------------------
# 内置列：key 即 tasks 核心列名，值存核心表；自定义列 key = c{id} 存 task_cells
BUILTIN_COLUMNS = [
    ("level1", "一级计划", "text", "", 0, 1),
    ("title", "任务", "text", "", 1, 0),
    ("progress", "状态", "select", '["未完成","进行中","已完成"]', 2, 0),
    ("completion", "完成度", "number", "", 3, 0),
    ("note", "备注", "text", "", 4, 0),
    ("start_date", "开始日期", "date", "", 5, 0),
    ("end_date", "结束日期", "date", "", 6, 0),
]

# 系统内置模板（仅重建自定义列部分，内置列始终保留）
PRESET_TEMPLATES = [
    ("标准计划表", [
        # 仅内置列的默认形态，无自定义列
    ]),
    ("项目跟踪", [
        {"label": "负责人", "ftype": "text", "options": []},
        {"label": "优先级", "ftype": "select", "options": ["高", "中", "低"]},
        {"label": "风险", "ftype": "select", "options": ["低", "中", "高"]},
    ]),
    ("周计划", [
        {"label": "负责人", "ftype": "text", "options": []},
        {"label": "标签", "ftype": "text", "options": []},
    ]),
]


def _seed_columns(conn: sqlite3.Connection) -> None:
    """幂等：首次为空时写入内置列与系统模板。"""
    n = conn.execute("SELECT COUNT(*) AS n FROM task_columns").fetchone()["n"]
    if n == 0:
        for key, label, ftype, opts, pos, pinned in BUILTIN_COLUMNS:
            conn.execute(
                "INSERT INTO task_columns (key, label, ftype, options, position, pinned, builtin, visible) "
                "VALUES (?,?,?,?,?,?,1,1)",
                (key, label, ftype, opts, pos, pinned),
            )
    # 模板种子（仅当模板表为空）
    tn = conn.execute("SELECT COUNT(*) AS n FROM task_templates").fetchone()["n"]
    if tn == 0:
        for name, cols in PRESET_TEMPLATES:
            conn.execute(
                "INSERT INTO task_templates (name, builtin, columns_json) VALUES (?, 1, ?)",
                (name, json.dumps(cols, ensure_ascii=False)),
            )


def list_columns() -> list[dict]:
    """所有列定义，按 position 升序。"""
    conn = _get_conn()
    try:
        _seed_columns(conn)
        conn.commit()
        rows = conn.execute("SELECT * FROM task_columns ORDER BY position, id").fetchall()
        return [_col_to_dict(r) for r in rows]
    finally:
        conn.close()


def _col_to_dict(r: sqlite3.Row) -> dict:
    opts = json.loads(r["options"]) if r["options"] else []
    return {
        "id": r["id"],
        "key": r["key"],
        "label": r["label"],
        "ftype": r["ftype"],
        "options": opts,
        "position": r["position"],
        "pinned": bool(r["pinned"]),
        "builtin": bool(r["builtin"]),
        "visible": bool(r["visible"]),
    }


def add_column(label: str, ftype: str, options: list[str] | None = None) -> dict:
    """新增自定义列（builtin=0），返回新列定义。"""
    conn = _get_conn()
    try:
        _seed_columns(conn)
        cur = conn.execute(
            "INSERT INTO task_columns (key, label, ftype, options, position, pinned, builtin, visible) "
            "VALUES ('', ?, ?, ?, "
            "(SELECT COALESCE(MAX(position),0)+1 FROM task_columns), 0, 0, 1)",
            (label, ftype, json.dumps(options or [], ensure_ascii=False)),
        )
        cid = cur.lastrowid
        conn.execute("UPDATE task_columns SET key=? WHERE id=?", (f"c{cid}", cid))
        conn.commit()
        return _col_to_dict(conn.execute("SELECT * FROM task_columns WHERE id=?", (cid,)).fetchone())
    finally:
        conn.close()


def update_column(
    cid: int,
    label: str | None = None,
    ftype: str | None = None,
    options: list[str] | None = None,
    visible: bool | None = None,
    pinned: bool | None = None,
) -> dict | None:
    """修改列（重命名 / 改类型 / 选项 / 可见 / 固定）。内置列不允许改类型与删除。"""
    conn = _get_conn()
    try:
        row = conn.execute("SELECT * FROM task_columns WHERE id=?", (cid,)).fetchone()
        if row is None:
            return None
        sets, params = [], []
        if label is not None:
            sets.append("label=?")
            params.append(label)
        if ftype is not None and not row["builtin"]:
            sets.append("ftype=?")
            params.append(ftype)
        if options is not None and not row["builtin"]:
            sets.append("options=?")
            params.append(json.dumps(options, ensure_ascii=False))
        if visible is not None:
            sets.append("visible=?")
            params.append(1 if visible else 0)
        if pinned is not None and not row["builtin"]:
            sets.append("pinned=?")
            params.append(1 if pinned else 0)
        if sets:
            conn.execute(f"UPDATE task_columns SET {','.join(sets)} WHERE id=?", (*params, cid))
        conn.commit()
        return _col_to_dict(conn.execute("SELECT * FROM task_columns WHERE id=?", (cid,)).fetchone())
    finally:
        conn.close()


def delete_column(cid: int) -> bool:
    """删除自定义列（内置列拒绝）。"""
    conn = _get_conn()
    try:
        row = conn.execute("SELECT * FROM task_columns WHERE id=?", (cid,)).fetchone()
        if row is None or row["builtin"]:
            return False
        conn.execute("DELETE FROM task_cells WHERE col_key=?", (row["key"],))
        conn.execute("DELETE FROM task_columns WHERE id=?", (cid,))
        conn.commit()
        return True
    finally:
        conn.close()


def set_columns_order(ordered_ids: list[int]) -> None:
    """整体重排列（拖拽结果）。ordered_ids 为全部可见/隐藏列的完整顺序。"""
    conn = _get_conn()
    try:
        for i, cid in enumerate(ordered_ids):
            conn.execute("UPDATE task_columns SET position=? WHERE id=?", (i, cid))
        conn.commit()
    finally:
        conn.close()


def upsert_cells(task_id: int, fields: dict[str, object]) -> None:
    """批量写入自定义列的值（task_cells，EAV）。内置列 key 自动忽略。"""
    conn = _get_conn()
    try:
        builtin = {r["key"] for r in conn.execute("SELECT key FROM task_columns WHERE builtin=1")}
        for key, value in fields.items():
            if key in builtin:
                continue
            val = "" if value is None else str(value)
            conn.execute(
                "INSERT INTO task_cells (task_id, col_key, value) VALUES (?,?,?) "
                "ON CONFLICT(task_id, col_key) DO UPDATE SET value=excluded.value",
                (task_id, key, val),
            )
        conn.commit()
    finally:
        conn.close()


def list_templates() -> list[dict]:
    """系统内置模板列表。"""
    conn = _get_conn()
    try:
        rows = conn.execute("SELECT * FROM task_templates ORDER BY id").fetchall()
        return [{"id": r["id"], "name": r["name"], "columns": json.loads(r["columns_json"])} for r in rows]
    finally:
        conn.close()


def apply_columns_json(cols: list) -> list[dict]:
    """按给定的列定义重建自定义列：清空旧自定义列及其单元格后重建，返回新列定义。

    供统一模板库（E3）套用「任务表」模板使用。
    """
    conn = _get_conn()
    try:
        _seed_columns(conn)
        # 删除所有自定义列与对应单元格
        custom = conn.execute("SELECT key FROM task_columns WHERE builtin=0").fetchall()
        for c in custom:
            conn.execute("DELETE FROM task_cells WHERE col_key=?", (c["key"],))
        conn.execute("DELETE FROM task_columns WHERE builtin=0")
        base = conn.execute("SELECT COALESCE(MAX(position),0) FROM task_columns").fetchone()[0]
        for i, c in enumerate(cols):
            conn.execute(
                "INSERT INTO task_columns (key, label, ftype, options, position, pinned, builtin, visible) "
                "VALUES ('', ?, ?, ?, ?, 0, 0, 1)",
                (c.get("label", ""), c.get("ftype", "text"), json.dumps(c.get("options", []), ensure_ascii=False), base + 1 + i),
            )
            cid = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
            conn.execute("UPDATE task_columns SET key=? WHERE id=?", (f"c{cid}", cid))
        conn.commit()
        return list_columns()
    finally:
        conn.close()


def apply_template(name: str) -> list[dict]:
    """套用预设模板：清空自定义列及其单元格，按模板重建自定义列，返回新列定义。"""
    conn = _get_conn()
    try:
        _seed_columns(conn)
        row = conn.execute("SELECT * FROM task_templates WHERE name=?", (name,)).fetchone()
        if row is None:
            return list_columns()
        cols = json.loads(row["columns_json"])
    finally:
        conn.close()
    return apply_columns_json(cols)


# ---------------------------------------------------------------------------
# 字段规范化
# ---------------------------------------------------------------------------
def norm_progress(progress: str) -> str:
    """非法状态回落为默认「未完成」。"""
    return progress if progress in PROGRESS_OPTIONS else DEFAULT_PROGRESS


def norm_completion(value: object) -> int:
    """完成度规范到 0-100 的整数，非法值回落为 0。"""
    try:
        n = int(value)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return DEFAULT_COMPLETION
    return max(COMPLETION_MIN, min(COMPLETION_MAX, n))


def _row_to_dict(row: sqlite3.Row) -> dict:
    d = dict(row)
    d["progress"] = norm_progress(d.get("progress", ""))
    d["completion"] = norm_completion(d.get("completion"))
    return d


# ---------------------------------------------------------------------------
# CRUD
# ---------------------------------------------------------------------------
def create_task(
    level1: str,
    title: str,
    progress: str,
    completion: object,
    note: str,
    start_date: str,
    end_date: str,
) -> dict:
    conn = _get_conn()
    try:
        cur = conn.execute(
            "INSERT INTO tasks (level1, title, progress, completion, note, start_date, end_date) "
            "VALUES (?, ?, ?, ?, ?, ?, ?)",
            (
                level1 or "",
                title,
                norm_progress(progress),
                norm_completion(completion),
                note,
                start_date,
                end_date,
            ),
        )
        conn.commit()
        row = conn.execute("SELECT * FROM tasks WHERE id = ?", (cur.lastrowid,)).fetchone()
        return _row_to_dict(row)
    finally:
        conn.close()


def list_tasks(month: Optional[str] = None) -> list[dict]:
    """列出任务，按结束日期升序排序（无结束日期的排最后）。

    month 为 'YYYY-MM'，仅返回在该月内活动的任务；为空则返回全部。
    每个任务额外带 `fields`：自定义列（task_cells）的 {col_key: value}。
    """
    conn = _get_conn()
    try:
        if month:
            first_day = f"{month}-01"
            y, m = int(month[:4]), int(month[5:7])
            ny, nm = (y + 1, 1) if m == 12 else (y, m + 1)
            next_month = f"{ny:04d}-{nm:02d}-01"
            rows = conn.execute(
                "SELECT * FROM tasks WHERE "
                "(start_date = '' OR start_date < ?) AND "
                "(end_date = '' OR end_date >= ?) " + ORDER_BY_SQL,
                (next_month, first_day),
            ).fetchall()
        else:
            rows = conn.execute(f"SELECT * FROM tasks {ORDER_BY_SQL}").fetchall()
        tasks = [_row_to_dict(r) for r in rows]
        if tasks:
            ids = [t["id"] for t in tasks]
            placeholders = ",".join("?" * len(ids))
            cells = conn.execute(
                f"SELECT task_id, col_key, value FROM task_cells WHERE task_id IN ({placeholders})",
                ids,
            ).fetchall()
            by_task: dict[int, dict] = {}
            for c in cells:
                by_task.setdefault(c["task_id"], {})[c["col_key"]] = c["value"]
            for t in tasks:
                t["fields"] = by_task.get(t["id"], {})
        else:
            for t in tasks:
                t["fields"] = {}
        return tasks
    finally:
        conn.close()


def get_task(tid: int) -> Optional[dict]:
    conn = _get_conn()
    try:
        row = conn.execute("SELECT * FROM tasks WHERE id = ?", (tid,)).fetchone()
        return _row_to_dict(row) if row is not None else None
    finally:
        conn.close()


def update_task(
    tid: int,
    level1: str,
    title: str,
    progress: str,
    completion: object,
    note: str,
    start_date: str,
    end_date: str,
) -> Optional[dict]:
    conn = _get_conn()
    try:
        cur = conn.execute(
            "UPDATE tasks SET level1=?, title=?, progress=?, completion=?, note=?, "
            "start_date=?, end_date=?, updated_at=datetime('now','localtime') "
            "WHERE id=?",
            (
                level1 or "",
                title,
                norm_progress(progress),
                norm_completion(completion),
                note,
                start_date,
                end_date,
                tid,
            ),
        )
        conn.commit()
        if cur.rowcount == 0:
            return None
        row = conn.execute("SELECT * FROM tasks WHERE id = ?", (tid,)).fetchone()
        return _row_to_dict(row)
    finally:
        conn.close()


def delete_task(tid: int) -> bool:
    conn = _get_conn()
    try:
        cur = conn.execute("DELETE FROM tasks WHERE id = ?", (tid,))
        conn.commit()
        return cur.rowcount > 0
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# 统计
# ---------------------------------------------------------------------------
def stats(month: Optional[str] = None) -> dict:
    """按用户手动设定的状态统计数量。"""
    tasks = list_tasks(month=month)
    counts = {p: 0 for p in (PROGRESS_DONE, PROGRESS_TODO, PROGRESS_DOING)}
    for t in tasks:
        counts[t["progress"]] = counts.get(t["progress"], 0) + 1
    return {
        "total": len(tasks),
        "done": counts[PROGRESS_DONE],
        "todo": counts[PROGRESS_TODO],
        "doing": counts[PROGRESS_DOING],
    }


# ---------------------------------------------------------------------------
# Excel 导出
# ---------------------------------------------------------------------------
def _task_month(t: dict) -> str:
    """任务归属月份（YYYY-MM），优先用结束日期，无日期归为「未排期」。"""
    d = t.get("end_date") or t.get("start_date") or ""
    return d[:7] if len(d) >= 7 else "未排期"


def export_excel() -> bytes:
    """导出任务为 Excel，按月份分 sheet，极简网格、无状态配色。

    列跟随当前列定义（内置列 + 自定义列）动态生成。
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise RuntimeError("缺少 openpyxl，请先安装：pip install openpyxl")

    cols = list_columns()
    visible_cols = [c for c in cols if c["visible"]]
    # 内置列取核心字段；自定义列取 fields[col_key]
    builtin_keys = {"level1", "title", "progress", "completion", "note", "start_date", "end_date"}

    tasks = list_tasks()

    by_month: dict[str, list[dict]] = {}
    for t in tasks:
        by_month.setdefault(_task_month(t), []).append(t)
    ordered_months = sorted(by_month.keys(), key=lambda m: (m != "未排期", m))

    wb = Workbook()
    wb.remove(wb.active)

    headers = [c["label"] for c in visible_cols]
    widths = [22] * len(visible_cols)
    # 备注列宽一些
    for i, c in enumerate(visible_cols):
        if c["key"] == "note":
            widths[i] = 40

    thin = Side(style="thin", color="D9D5CC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def cell_value(t: dict, c: dict) -> object:
        if c["key"] in builtin_keys:
            if c["key"] == "completion":
                return f"{t.get('completion', 0)}%"
            return t.get(c["key"], "")
        return t.get("fields", {}).get(c["key"], "")

    def fill_sheet(ws, month_label: str, month_tasks: list[dict]):
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        title_cell = ws.cell(row=1, column=1, value=f"MaWork 日程 · {month_label}")
        title_cell.font = Font(size=14, bold=True)
        title_cell.alignment = Alignment(horizontal="center")
        ws.row_dimensions[1].height = 24

        header_fill = PatternFill("solid", fgColor="F3F1EC")
        for col, h in enumerate(headers, start=1):
            c = ws.cell(row=2, column=col, value=h)
            c.font = Font(bold=True)
            c.fill = header_fill
            c.border = border
            c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 20

        for i, t in enumerate(month_tasks, start=3):
            values = [cell_value(t, c) for c in visible_cols]
            for col, v in enumerate(values, start=1):
                c = ws.cell(row=i, column=col, value=v)
                c.border = border
                is_note = visible_cols[col - 1]["key"] == "note"
                c.alignment = Alignment(vertical="center", wrap_text=is_note)
            note_lines = min((t.get("note") or "").count("\n") + 1, 14)
            if any(c["key"] == "note" for c in visible_cols) and note_lines > 1:
                ws.row_dimensions[i].height = 18 * note_lines

        for col, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(col)].width = w
        ws.freeze_panes = "A3"

    for month in ordered_months:
        ws = wb.create_sheet(title=str(month))
        fill_sheet(ws, str(month), by_month[month])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
